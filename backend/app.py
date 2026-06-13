from flask import Flask, request, jsonify, send_from_directory, abort, make_response
from flask_cors import CORS
import sqlite3, os, uuid, datetime, io
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'issues.db')
UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_EXTENSIONS = {
    'png','jpg','jpeg','gif','webp','bmp',
    'pdf','txt','log','md','csv','json',
    'zip','tar','gz','doc','docx','xls','xlsx'
}

STATIC_DIR = os.path.join(BASE_DIR, 'static')
app = Flask(__name__, static_folder=STATIC_DIR, static_url_path='')
CORS(app)
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 32 MB


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.executescript('''
            CREATE TABLE IF NOT EXISTS issues (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                title       TEXT NOT NULL,
                severity    TEXT NOT NULL DEFAULT 'medium',
                status      TEXT NOT NULL DEFAULT 'open',
                project     TEXT,
                steps       TEXT,
                notes       TEXT,
                created_at  TEXT NOT NULL,
                updated_at  TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS attachments (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                issue_id    INTEGER NOT NULL REFERENCES issues(id) ON DELETE CASCADE,
                filename    TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                mime_type   TEXT,
                size        INTEGER,
                uploaded_at TEXT NOT NULL
            );
        ''')


init_db()


def now():
    return datetime.datetime.utcnow().isoformat() + 'Z'


def row_to_dict(row):
    return dict(row)


def issue_with_attachments(conn, issue_id):
    issue = conn.execute('SELECT * FROM issues WHERE id=?', (issue_id,)).fetchone()
    if not issue:
        return None
    d = row_to_dict(issue)
    atts = conn.execute(
        'SELECT id, filename, stored_name, mime_type, size, uploaded_at FROM attachments WHERE issue_id=? ORDER BY uploaded_at',
        (issue_id,)
    ).fetchall()
    d['attachments'] = [row_to_dict(a) for a in atts]
    return d


# ── Issues ────────────────────────────────────────────────────────────────────

@app.route('/api/issues', methods=['GET'])
def list_issues():
    status  = request.args.get('status')
    project = request.args.get('project')
    severity= request.args.get('severity')
    q       = request.args.get('q', '').strip()

    sql = 'SELECT * FROM issues WHERE 1=1'
    params = []
    if status:   sql += ' AND status=?';   params.append(status)
    if project:  sql += ' AND project=?';  params.append(project)
    if severity: sql += ' AND severity=?'; params.append(severity)
    if q:        sql += ' AND (title LIKE ? OR notes LIKE ? OR steps LIKE ?)'; params += [f'%{q}%']*3
    sql += ' ORDER BY id DESC'

    with get_db() as conn:
        rows = conn.execute(sql, params).fetchall()
        issues = []
        for row in rows:
            d = row_to_dict(row)
            count = conn.execute('SELECT COUNT(*) FROM attachments WHERE issue_id=?', (d['id'],)).fetchone()[0]
            d['attachment_count'] = count
            issues.append(d)
    return jsonify(issues)


@app.route('/api/issues', methods=['POST'])
def create_issue():
    data = request.get_json()
    t = now()
    with get_db() as conn:
        cur = conn.execute(
            'INSERT INTO issues (title,severity,status,project,steps,notes,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)',
            (data['title'], data.get('severity','medium'), data.get('status','open'),
             data.get('project'), data.get('steps'), data.get('notes'), t, t)
        )
        conn.commit()
        issue = issue_with_attachments(conn, cur.lastrowid)
    return jsonify(issue), 201


@app.route('/api/issues/<int:issue_id>', methods=['GET'])
def get_issue(issue_id):
    with get_db() as conn:
        issue = issue_with_attachments(conn, issue_id)
    if not issue:
        abort(404)
    return jsonify(issue)


@app.route('/api/issues/<int:issue_id>', methods=['PUT'])
def update_issue(issue_id):
    data = request.get_json()
    t = now()
    with get_db() as conn:
        conn.execute(
            '''UPDATE issues SET title=?,severity=?,status=?,project=?,steps=?,notes=?,updated_at=?
               WHERE id=?''',
            (data['title'], data.get('severity','medium'), data.get('status','open'),
             data.get('project'), data.get('steps'), data.get('notes'), t, issue_id)
        )
        conn.commit()
        issue = issue_with_attachments(conn, issue_id)
    if not issue:
        abort(404)
    return jsonify(issue)


@app.route('/api/issues/<int:issue_id>', methods=['DELETE'])
def delete_issue(issue_id):
    with get_db() as conn:
        atts = conn.execute('SELECT stored_name FROM attachments WHERE issue_id=?', (issue_id,)).fetchall()
        for a in atts:
            path = os.path.join(UPLOAD_DIR, a['stored_name'])
            if os.path.exists(path):
                os.remove(path)
        conn.execute('DELETE FROM issues WHERE id=?', (issue_id,))
        conn.commit()
    return '', 204


# ── Attachments ───────────────────────────────────────────────────────────────

@app.route('/api/issues/<int:issue_id>/attachments', methods=['POST'])
def upload_attachment(issue_id):
    with get_db() as conn:
        if not conn.execute('SELECT 1 FROM issues WHERE id=?', (issue_id,)).fetchone():
            abort(404)

    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    f = request.files['file']
    if not f.filename:
        return jsonify({'error': 'Empty filename'}), 400
    if not allowed_file(f.filename):
        return jsonify({'error': 'File type not allowed'}), 400

    ext = f.filename.rsplit('.', 1)[1].lower()
    stored = f'{uuid.uuid4().hex}.{ext}'
    f.save(os.path.join(UPLOAD_DIR, stored))
    size = os.path.getsize(os.path.join(UPLOAD_DIR, stored))

    with get_db() as conn:
        cur = conn.execute(
            'INSERT INTO attachments (issue_id,filename,stored_name,mime_type,size,uploaded_at) VALUES (?,?,?,?,?,?)',
            (issue_id, secure_filename(f.filename), stored, f.content_type, size, now())
        )
        conn.commit()
        att = row_to_dict(conn.execute('SELECT * FROM attachments WHERE id=?', (cur.lastrowid,)).fetchone())
    return jsonify(att), 201


@app.route('/api/attachments/<int:att_id>', methods=['DELETE'])
def delete_attachment(att_id):
    with get_db() as conn:
        att = conn.execute('SELECT * FROM attachments WHERE id=?', (att_id,)).fetchone()
        if not att:
            abort(404)
        path = os.path.join(UPLOAD_DIR, att['stored_name'])
        if os.path.exists(path):
            os.remove(path)
        conn.execute('DELETE FROM attachments WHERE id=?', (att_id,))
        conn.commit()
    return '', 204


@app.route('/api/attachments/<int:att_id>/download')
def download_attachment(att_id):
    with get_db() as conn:
        att = conn.execute('SELECT * FROM attachments WHERE id=?', (att_id,)).fetchone()
    if not att:
        abort(404)
    return send_from_directory(UPLOAD_DIR, att['stored_name'], as_attachment=True,
                               download_name=att['filename'])


@app.route('/api/attachments/<int:att_id>/view')
def view_attachment(att_id):
    with get_db() as conn:
        att = conn.execute('SELECT * FROM attachments WHERE id=?', (att_id,)).fetchone()
    if not att:
        abort(404)
    return send_from_directory(UPLOAD_DIR, att['stored_name'],
                               mimetype=att['mime_type'], as_attachment=False)


# ── Meta ──────────────────────────────────────────────────────────────────────

@app.route('/api/projects', methods=['GET'])
def list_projects():
    with get_db() as conn:
        rows = conn.execute(
            "SELECT DISTINCT project FROM issues WHERE project IS NOT NULL AND project != '' ORDER BY project"
        ).fetchall()
    return jsonify([r['project'] for r in rows])


@app.route('/api/stats', methods=['GET'])
def stats():
    with get_db() as conn:
        total  = conn.execute('SELECT COUNT(*) FROM issues').fetchone()[0]
        open_  = conn.execute("SELECT COUNT(*) FROM issues WHERE status='open'").fetchone()[0]
        inprog = conn.execute("SELECT COUNT(*) FROM issues WHERE status='in-progress'").fetchone()[0]
        closed = conn.execute("SELECT COUNT(*) FROM issues WHERE status='resolved'").fetchone()[0]
        crit   = conn.execute("SELECT COUNT(*) FROM issues WHERE severity='critical'").fetchone()[0]
    return jsonify({'total':total,'open':open_,'in_progress':inprog,'resolved':closed,'critical':crit})




@app.route('/api/export', methods=['GET'])
def export_issues():
    with get_db() as conn:
        issues = conn.execute('SELECT * FROM issues ORDER BY id').fetchall()
        issues = [dict(i) for i in issues]
        for issue in issues:
            atts = conn.execute(
                'SELECT filename, size, uploaded_at FROM attachments WHERE issue_id=?',
                (issue['id'],)
            ).fetchall()
            issue['attachments'] = [dict(a) for a in atts]

    wb = openpyxl.Workbook()

    # ── colour palette ────────────────────────────────────────────────
    SEV_FILLS = {
        'critical': PatternFill('solid', fgColor='3D1A1A'),
        'high':     PatternFill('solid', fgColor='3D2A10'),
        'medium':   PatternFill('solid', fgColor='3D3510'),
        'low':      PatternFill('solid', fgColor='0F2535'),
    }
    SEV_FONTS = {
        'critical': Font(color='F0625A', bold=True),
        'high':     Font(color='F5944A', bold=True),
        'medium':   Font(color='F5C842', bold=True),
        'low':      Font(color='5AB4F5', bold=True),
    }
    STATUS_FILLS = {
        'open':        PatternFill('solid', fgColor='2A1010'),
        'in-progress': PatternFill('solid', fgColor='2A2510'),
        'resolved':    PatternFill('solid', fgColor='0F2A1A'),
    }
    STATUS_FONTS = {
        'open':        Font(color='F0625A', bold=True),
        'in-progress': Font(color='F5C842', bold=True),
        'resolved':    Font(color='3DD68C', bold=True),
    }
    HEADER_FILL = PatternFill('solid', fgColor='1C1C20')
    HEADER_FONT = Font(color='7C6AF7', bold=True)
    ROW_FILL_A  = PatternFill('solid', fgColor='141417')
    ROW_FILL_B  = PatternFill('solid', fgColor='0D0D0F')
    TEXT_FONT   = Font(color='B0B0BC')
    WRAP        = Alignment(wrap_text=True, vertical='top')
    thin        = Side(style='thin', color='2E2E33')
    BORDER      = Border(bottom=thin)

    def style_header(cell, text):
        cell.value = text
        cell.font  = HEADER_FONT
        cell.fill  = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = Border(bottom=Side(style='medium', color='7C6AF7'))

    def style_cell(cell, row_idx, wrap=False):
        cell.fill   = ROW_FILL_A if row_idx % 2 == 0 else ROW_FILL_B
        cell.font   = TEXT_FONT
        cell.border = BORDER
        cell.alignment = WRAP if wrap else Alignment(vertical='top')

    # ── Issues sheet ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Issues'
    ws.sheet_view.showGridLines = False

    headers = ['ID', 'Title', 'Severity', 'Status', 'Project', 'Steps to Reproduce', 'Notes', 'Attachments', 'Created', 'Updated']
    col_widths = [6, 40, 10, 12, 16, 40, 40, 20, 20, 20]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        style_header(ws.cell(row=1, column=col), h)
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    for row_idx, issue in enumerate(issues, 2):
        att_names = ', '.join(a['filename'] for a in issue['attachments']) if issue['attachments'] else ''
        created = issue['created_at'][:19].replace('T', ' ') if issue['created_at'] else ''
        updated = issue['updated_at'][:19].replace('T', ' ') if issue['updated_at'] else ''

        values = [
            issue['id'], issue['title'], issue['severity'], issue['status'],
            issue['project'] or '', issue['steps'] or '', issue['notes'] or '',
            att_names, created, updated
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            wrap = col in (2, 6, 7, 8)
            style_cell(cell, row_idx, wrap=wrap)

            if col == 3:  # severity
                cell.fill = SEV_FILLS.get(issue['severity'], ROW_FILL_A)
                cell.font = SEV_FONTS.get(issue['severity'], TEXT_FONT)
                cell.alignment = Alignment(horizontal='center', vertical='top')
            elif col == 4:  # status
                cell.fill = STATUS_FILLS.get(issue['status'], ROW_FILL_A)
                cell.font = STATUS_FONTS.get(issue['status'], TEXT_FONT)
                cell.alignment = Alignment(horizontal='center', vertical='top')

        ws.row_dimensions[row_idx].height = 60

    # ── Attachments sheet ─────────────────────────────────────────────
    ws2 = wb.create_sheet('Attachments')
    ws2.sheet_view.showGridLines = False

    att_headers = ['Issue ID', 'Issue Title', 'Filename', 'Size (KB)', 'Uploaded']
    att_widths   = [10, 40, 40, 12, 20]

    for col, (h, w) in enumerate(zip(att_headers, att_widths), 1):
        style_header(ws2.cell(row=1, column=col), h)
        ws2.column_dimensions[get_column_letter(col)].width = w

    ws2.row_dimensions[1].height = 22
    ws2.freeze_panes = 'A2'

    att_row = 2
    for issue in issues:
        for att in issue['attachments']:
            size_kb = round(att['size'] / 1024, 1) if att['size'] else 0
            uploaded = att['uploaded_at'][:19].replace('T', ' ') if att['uploaded_at'] else ''
            vals = [issue['id'], issue['title'], att['filename'], size_kb, uploaded]
            for col, val in enumerate(vals, 1):
                cell = ws2.cell(row=att_row, column=col, value=val)
                style_cell(cell, att_row)
            ws2.row_dimensions[att_row].height = 18
            att_row += 1

    # ── Stream response ───────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"issues-export-{datetime.date.today().isoformat()}.xlsx"
    response = make_response(buf.read())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@app.route('/', defaults={'path': ''})
@app.route('/<path:path>')
def serve_frontend(path):
    if path and os.path.exists(os.path.join(STATIC_DIR, path)):
        return send_from_directory(STATIC_DIR, path)
    return send_from_directory(STATIC_DIR, 'index.html')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5050, debug=False)
